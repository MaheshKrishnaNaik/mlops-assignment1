"""Build docs/MLOps_Assignment1_Report.docx from project artifacts. Run: python scripts/make_report.py"""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "screenshots"
REPO = "https://github.com/MaheshKrishnaNaik/mlops-assignment1"
M = json.loads((ROOT / "models" / "metrics.json").read_text())["results"]
doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def para(text, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if align:
        p.alignment = align
    return p


def bullets(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)


def table(rows, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]), style="Light Grid Accent 1")
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            t.cell(i, j).text = str(v)
            if i == 0 and header:
                for r in t.cell(i, j).paragraphs[0].runs:
                    r.bold = True
    doc.add_paragraph()


def figure(path, caption, width=5.8):
    p = Path(path)
    if p.exists():
        doc.add_picture(str(p), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        para(f"[Insert screenshot: {p.name}]", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(caption, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)


def h(text, lvl=1):
    doc.add_heading(text, lvl)


# ---------------- Title ----------------
for _ in range(6):
    doc.add_paragraph()
para("Machine Learning Operations (MLOps) - AIMLCZG523", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Assignment 01", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para("End-to-End ML Model Development, CI/CD and Production Deployment", bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("Heart Disease Risk Prediction (UCI Heart Disease Dataset)", align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(3):
    doc.add_paragraph()
table([["Student ID", "2025af05162"],
       ["Email", "2025af05162@wilp.bits-pilani.ac.in"],
       ["GitHub repository", REPO],
       ["Deployed API", "Local Kubernetes (Minikube) - access instructions in Section 8"]], header=False)
doc.add_page_break()

# ---------------- 1 Overview ----------------
h("1. Introduction and Architecture")
para("The goal of this assignment is to build a classifier that predicts the presence of heart disease from "
     "patient health data, and to ship it the way a production ML service is shipped: reproducible training, "
     "tracked experiments, automated tests and CI, a containerised API, a Kubernetes deployment and basic "
     "monitoring. Every stage is automated by a script in the repository so that the whole flow can be re-run "
     "from a clean checkout using only requirements.txt.")
h("1.1 Architecture", 2)
code("UCI data --> src/data.py (clean) --> src/train.py (LogReg + RandomForest, GridSearchCV, 5-fold CV)\n"
     "     --> MLflow tracking (params, metrics, plots, models)\n"
     "     --> models/model.joblib (one sklearn Pipeline: preprocessing + estimator)\n"
     "     --> FastAPI service (/predict, /health, /metrics) --> Docker image\n"
     "     --> Kubernetes Deployment + LoadBalancer Service (Minikube)\n"
     "     --> Prometheus scrapes /metrics --> Grafana dashboard\n"
     "GitHub Actions: flake8 -> pytest -> data -> train -> upload artifacts -> docker build + smoke test")
h("1.2 Repository layout", 2)
table([["Path", "Purpose"],
       ["src/", "data.py (download/clean), features.py (pipeline), train.py (training + MLflow), app.py (API)"],
       ["tests/", "pytest unit tests for data cleaning, model pipeline and API"],
       ["notebooks/", "eda.py (EDA), inference.py (load model and predict)"],
       ["data/", "raw and cleaned dataset"],
       ["models/", "model.joblib and metrics.json"],
       ["k8s/, monitoring/", "Deployment/Service/Ingress manifests; Prometheus + Grafana compose file"],
       [".github/workflows/ci.yml", "CI pipeline"],
       ["screenshots/", "EDA plots and run evidence used in this report"]])

# ---------------- 2 Setup ----------------
h("2. Setup and Installation")
code("git clone " + REPO + ".git\ncd mlops-assignment1\npython3 -m venv .venv && source .venv/bin/activate\n"
     "pip install -r requirements.txt\n"
     "python -m src.data            # download + clean data\n"
     "PYTHONPATH=. python notebooks/eda.py\n"
     "python -m src.train           # train, track, save model\n"
     "pytest -v && flake8 src tests\n"
     "mlflow ui                     # http://127.0.0.1:5000")
para("The dataset is fetched by src/data.py from the UCI repository "
     "(https://archive.ics.uci.edu/static/public/45/data.csv), so no manual download is required. The raw and "
     "cleaned CSV files are also committed for convenience. All dependencies are listed in requirements.txt with "
     "version ranges; Python 3.12 was used.")

# ---------------- 3 Data & EDA ----------------
h("3. Data Acquisition and Exploratory Data Analysis")
h("3.1 Dataset", 2)
para("The Heart Disease UCI dataset (Cleveland subset) has 303 patients and 13 predictors: age, sex, chest pain type "
     "(cp), resting blood pressure (trestbps), cholesterol (chol), fasting blood sugar (fbs), resting ECG "
     "(restecg), maximum heart rate (thalach), exercise-induced angina (exang), ST depression (oldpeak), slope "
     "of the ST segment (slope), number of major vessels coloured by fluoroscopy (ca) and thalassemia (thal). "
     "The original target 'num' ranges from 0 to 4; following the assignment it is binarised so that any value "
     "greater than 0 means heart disease is present.")
h("3.2 Cleaning", 2)
bullets(["Missing values appear only in 'ca' (4 rows) and 'thal' (2 rows). Both are categorical codes, so they are "
         "imputed with the column mode. Numeric columns would use the median.",
         "No duplicate rows were found in the raw data; the cleaning function still drops duplicates.",
         "All columns are coerced to numeric so that '?' placeholders become NaN before imputation.",
         "Categorical columns (sex, cp, fbs, restecg, exang, slope, ca, thal) are kept as integer codes in the "
         "cleaned file and one-hot encoded inside the modelling pipeline, which avoids leaking encodings "
         "between train and test splits."])
h("3.3 Findings", 2)
figure(SHOTS / "eda" / "class_balance.png", "Figure 1: Class balance (54.1% no disease, 45.9% disease).", 4.2)
para("The classes are close to balanced, so accuracy is a meaningful metric and no resampling is needed. "
     "Recall is still tracked because missing a sick patient is the costlier error in this domain.")
figure(SHOTS / "eda" / "histograms.png", "Figure 2: Distributions of numeric features.")
para("Age is roughly bell-shaped around 55; cholesterol is right-skewed with a few very high values; oldpeak is "
     "strongly right-skewed with many zeros. Standard scaling is therefore applied to the numeric features.")
figure(SHOTS / "eda" / "correlation_heatmap.png", "Figure 3: Correlation heatmap including the target.", 5.2)
para("The strongest correlations with the target are thal (0.52), ca (0.46), exang (0.43), oldpeak (0.42), "
     "cp (0.41) and thalach (-0.42, a lower maximum heart rate goes with disease). Fasting blood sugar (0.03) and "
     "cholesterol (0.09) are only weakly related. Patients with disease average 56.6 years against 52.6 and a "
     "maximum heart rate of 139 against 158 bpm.")
figure(SHOTS / "eda" / "categorical_vs_target.png", "Figure 4: Chest pain type, thal and ca versus the target.")

# ---------------- 4 Modelling ----------------
h("4. Feature Engineering and Model Development")
para("Preprocessing and the estimator form a single scikit-learn Pipeline (ColumnTransformer: StandardScaler for "
     "numeric columns, OneHotEncoder with handle_unknown='ignore' for categorical columns). Persisting the whole "
     "pipeline guarantees that the API applies exactly the transforms used in training, and unseen category "
     "values do not crash inference (covered by a unit test).")
para("The data was split 80/20 with stratification (random_state=42). Hyperparameters were tuned on the training "
     "part with GridSearchCV (stratified 5-fold, ROC-AUC scoring):")
table([["Model", "Search grid"],
       ["Logistic Regression", "C in {0.01, 0.1, 1, 10}, max_iter=2000"],
       ["Random Forest", "n_estimators in {100, 300}, max_depth in {3, 5, None}"]])
para("The tuned pipeline of each model was then cross-validated (5-fold) for accuracy, precision, recall and "
     "ROC-AUC, and finally evaluated once on the hold-out test set.")
lr, rf = M["logistic_regression"], M["random_forest"]


def row(name, r):
    return [name] + [f"{r[k]:.3f}" for k in ("cv_accuracy", "cv_precision", "cv_recall", "cv_roc_auc",
                                              "test_accuracy", "test_precision", "test_recall", "test_roc_auc")]


table([["Model", "CV Acc", "CV Prec", "CV Rec", "CV AUC", "Test Acc", "Test Prec", "Test Rec", "Test AUC"],
       row("Logistic Regression (C=1)", lr), row("Random Forest (100, depth=None)", rf)])
para("Model selection: Logistic Regression has the higher cross-validated accuracy, precision and ROC-AUC, is much "
     "simpler and produces well-behaved probabilities, which suits a risk-score API. Random Forest did not "
     "outperform it on this small dataset. The selection criterion is CV ROC-AUC (0.902 versus 0.897). The test "
     "set contains only 61 patients, so the test figures are noisy and were not used to choose the model.")

# ---------------- 5 Tracking ----------------
h("5. Experiment Tracking with MLflow")
para("Each model trains inside an MLflow run in the experiment 'heart-disease'. Per run the code logs: the best "
     "hyperparameters, all cross-validation metrics (cv_accuracy, cv_precision, cv_recall, cv_roc_auc), the "
     "hold-out metrics, the ROC curve and confusion matrix as image artifacts, and the fitted pipeline as an "
     "MLflow model. The MLflow UI (mlflow ui) is used to compare runs.")
figure(SHOTS / "mlflow_runs.png", "Figure 5: MLflow experiment with both runs (add screenshot).")
figure(SHOTS / "mlflow_compare.png", "Figure 6: MLflow run comparison (add screenshot).")

# ---------------- 6 Packaging ----------------
h("6. Model Packaging and Reproducibility")
bullets(["Final model saved as models/model.joblib (full Pipeline) and as an MLflow model artifact.",
         "requirements.txt pins compatible version ranges; the API image installs only the runtime subset.",
         "Fixed random_state=42 for the split, cross-validation and Random Forest.",
         "models/metrics.json stores the selected model name and all metrics for audit.",
         "notebooks/inference.py shows loading the saved pipeline and predicting for a new patient."])

# ---------------- 7 CI/CD ----------------
h("7. CI/CD Pipeline and Automated Testing")
h("7.1 Unit tests", 2)
table([["File", "What it verifies"],
       ["tests/test_data.py", "target binarisation, no nulls after cleaning, duplicate removal, '?' handling"],
       ["tests/test_model.py", "both pipelines fit/predict, probabilities sum to 1, unseen categories, metric ranges"],
       ["tests/test_api.py", "/health, /predict output, input validation (422), /metrics counters"]])
para("11 tests pass locally. One test exposed a genuine bug (an entirely missing categorical column crashed mode "
     "imputation); it was fixed in src/data.py.")
h("7.2 GitHub Actions", 2)
para("The workflow .github/workflows/ci.yml runs on every push and pull request: install requirements, flake8, "
     "pytest (JUnit XML), data preparation, model training, upload of test results, model and MLflow runs as "
     "artifacts, Docker build, and a smoke test that starts the container and calls /health and /predict. "
     "Because each step runs without continue-on-error, any lint failure, failing test or training error fails the "
     "workflow, and the logs show the failing step.")
figure(SHOTS / "github_actions_run.png", "Figure 7: Successful GitHub Actions run (add screenshot).")
figure(SHOTS / "github_actions_artifacts.png", "Figure 8: Uploaded workflow artifacts (add screenshot).")

# ---------------- 8 Docker / K8s ----------------
h("8. Containerisation and Deployment")
h("8.1 Docker", 2)
para("The API is a FastAPI application. POST /predict accepts the 13 features as JSON (validated with Pydantic "
     "ranges) and returns the predicted class, the confidence in that class and the disease probability.")
code("docker build -t heart-api:latest .\ndocker run -d --name heart-api -p 8001:8000 heart-api:latest\n"
     "curl -X POST http://localhost:8001/predict -H \"Content-Type: application/json\" \\\n"
     "  -d '{\"age\":63,\"sex\":1,\"cp\":1,\"trestbps\":145,\"chol\":233,\"fbs\":1,\"restecg\":2,\"thalach\":150,"
     "\"exang\":0,\"oldpeak\":2.3,\"slope\":3,\"ca\":0,\"thal\":6}'\n"
     "# {\"prediction\":0,\"confidence\":0.83,\"risk_probability\":0.17}")
figure(SHOTS / "docker_run.png", "Figure 9: Docker build/run and sample /predict call (add screenshot).")
h("8.2 Kubernetes", 2)
para("k8s/deployment.yaml defines a Deployment with 2 replicas, readiness and liveness probes on /health, resource "
     "requests/limits, and a Service of type LoadBalancer (port 80 to container port 8000). k8s/ingress.yaml "
     "optionally exposes it through an Ingress. It was deployed on Minikube:")
code("minikube start --driver=docker\ndocker build -t heart-api:latest .\nminikube image load heart-api:latest\n"
     "kubectl apply -f k8s/deployment.yaml\nkubectl get pods,svc\nkubectl port-forward svc/heart-api 8001:80\n"
     "curl http://localhost:8001/health")
para("Access instructions for local testing: after the commands above the API is at http://localhost:8001 "
     "(interactive docs at /docs).")
figure(SHOTS / "kubectl_pods_svc.png", "Figure 10: kubectl get pods,svc (add screenshot).")
figure(SHOTS / "k8s_predict.png", "Figure 11: /predict served through Kubernetes (add screenshot).")

# ---------------- 9 Monitoring ----------------
h("9. Monitoring and Logging")
para("Every HTTP request is logged with method, path, status and latency by a FastAPI middleware, and each "
     "prediction logs its class and confidence. The /metrics endpoint exposes Prometheus metrics: "
     "api_requests_total{path,status}, api_request_latency_seconds (histogram) and "
     "api_predictions_total{label}. monitoring/docker-compose.yml starts Prometheus (port 9090) and Grafana "
     "(port 3000).")
table([["Panel", "PromQL"],
       ["Request rate", "sum(rate(api_requests_total[1m])) by (path)"],
       ["p95 latency", "histogram_quantile(0.95, sum(rate(api_request_latency_seconds_bucket[1m])) by (le))"],
       ["Predictions by class", "api_predictions_total"]])
figure(SHOTS / "grafana_dashboard.png", "Figure 12: Grafana/Prometheus dashboard (add screenshot).")

# ---------------- 10 Conclusion ----------------
h("10. Conclusion and Limitations")
para("The project delivers a reproducible pipeline from raw UCI data to a monitored, containerised API on "
     "Kubernetes, with experiment tracking and CI that fails on lint, test or training errors.")
bullets(["The dataset is small (303 rows), so metrics have wide confidence intervals; results are indicative, "
         "not clinical.",
         "Only two model families were compared; boosting models or calibration could be explored.",
         "The service has no authentication and no model-drift monitoring; both would be needed in production.",
         "The model is baked into the image; a registry-based deployment (MLflow Model Registry) would allow "
         "updates without rebuilding."])
h("References", 2)
bullets(["UCI Machine Learning Repository - Heart Disease dataset.",
         "scikit-learn, MLflow, FastAPI, Prometheus and Kubernetes documentation.",
         "Code repository: " + REPO])

out = ROOT / "docs" / "MLOps_Assignment1_Report.docx"
out.parent.mkdir(exist_ok=True)
doc.save(out)
print("Saved", out)
